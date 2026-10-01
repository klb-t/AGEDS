"""Read-only, bounded inventory and source observations.

This module has no database, network, subprocess or model dependencies. A path,
filename or spreadsheet field is an observation, never an identity decision.
Original bytes remain the authority; decoded cells are an explicitly limited view.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import re
import stat
import sys
import tempfile
import unicodedata
import zipfile
from dataclasses import asdict, dataclass
from datetime import date, datetime, time, timezone
from pathlib import Path
from typing import Any

from .wav_header import MAX_PREFIX_BYTES, probe_wav_header

SCHEMA_VERSION = "ageds.source-scan/v1"
SCANNER_VERSION = "1.1.0"
_AUDIO = {".wav", ".mp3", ".m4a", ".aac", ".ogg", ".opus", ".flac", ".amr", ".wma"}
_VIDEO = {".mp4", ".mov", ".mkv", ".avi", ".webm"}


@dataclass(frozen=True)
class ScanLimits:
    max_files: int = 5000
    max_directory_entries: int = 20000
    max_total_entries: int = 50000
    max_sheets_per_file: int = 64
    max_depth: int = 24
    max_hash_bytes: int = 512 * 1024 * 1024
    max_total_hash_bytes: int = 2 * 1024 * 1024 * 1024
    max_table_bytes: int = 16 * 1024 * 1024
    max_xlsx_expanded_bytes: int = 64 * 1024 * 1024
    max_rows_per_file: int = 2000
    max_total_rows: int = 20000
    max_columns: int = 128
    max_cell_characters: int = 4096
    max_total_cell_characters: int = 2_000_000
    max_observations: int = 50000
    max_candidate_links: int = 10000
    max_wav_header_bytes: int = MAX_PREFIX_BYTES
    max_total_wav_header_bytes: int = 8 * 1024 * 1024

    def __post_init__(self):
        if any(not isinstance(v, int) or isinstance(v, bool) or v < 1 for v in asdict(self).values()):
            raise ValueError("Every scan limit must be a positive integer")
        if self.max_wav_header_bytes > MAX_PREFIX_BYTES:
            raise ValueError("WAV header per-file limit must not exceed 65536 bytes")


def _id(*parts: Any) -> str:
    return hashlib.sha256(json.dumps(parts, ensure_ascii=False, separators=(",", ":")).encode()).hexdigest()


def _header(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or "").lower()).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", text)


_HEADERS = {
    "phone": {"phone", "phonenumber", "phoneoraddress", "number", "numer", "numertelefonu", "telefon", "address", "caller", "callee", "sender", "recipient", "from", "to", "numerabonenta", "numerdocelowy"},
    "timestamp": {"date", "datetime", "timestamp", "time", "data", "datapolaczenia", "dataczas", "czaspolaczenia", "start", "startedat", "tsstart"},
    "duration": {"duration", "durationseconds", "durationsec", "durationms", "czasrozmowy", "dlugosc", "dlugoscrozmowy", "seconds", "sekundy"},
    "file_reference": {"file", "filename", "recording", "recordingfile", "nagranie", "plik", "nazwapliku", "path", "filepath"},
    "contact_label": {"name", "contact", "contactname", "kontakt", "nazwa", "osoba"},
}


def _kind_for_header(value: Any) -> str | None:
    h = _header(value)
    return next((kind for kind, names in _HEADERS.items() if h in names), None)


def _phone(value: Any) -> str | None:
    if not isinstance(value, str):
        return None  # Numeric cells may already have lost leading zeros/precision.
    if "@" in value:
        return None
    cleaned = re.sub(r"[\s().-]", "", value)
    return cleaned if re.fullmatch(r"\+?\d{6,15}", cleaned) else None


def _datetime_view(value: Any) -> dict:
    result = {"normalized": None, "timezone": "unknown", "interpretation": "unparsed"}
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        return {"normalized": value.isoformat(), "timezone": "unknown", "interpretation": "calendar_date"}
    elif isinstance(value, str):
        try:
            parsed = datetime.fromisoformat(value.strip().replace("Z", "+00:00"))
        except ValueError:
            return result
    else:
        return result
    return {"normalized": parsed.isoformat(), "timezone": str(parsed.utcoffset()) if parsed.tzinfo and parsed.utcoffset() is not None else "unknown", "interpretation": "iso_or_typed_datetime"}


def _duration_view(value: Any, header: Any = None) -> dict:
    result = {"seconds": None, "unit": "unknown"}
    if isinstance(value, str) and re.fullmatch(r"\d+:\d{2}:\d{2}", value.strip()):
        h, m, s = map(int, value.strip().split(":"))
        if m < 60 and s < 60:
            return {"seconds": h * 3600 + m * 60 + s, "unit": "hh:mm:ss"}
    unit = "milliseconds" if _header(header) == "durationms" else "seconds" if _header(header) in {"durationseconds", "durationsec", "seconds", "sekundy"} else "unknown"
    if unit != "unknown":
        try:
            number = float(str(value).replace(",", "."))
            if number >= 0 and number != float("inf") and number == number:
                return {"seconds": number / 1000 if unit == "milliseconds" else number, "unit": unit}
        except (ValueError, TypeError):
            pass
    return result


class _Scan:
    def __init__(self, root: Path, limits: ScanLimits):
        self.root, self.limits = root, limits
        self.hash_bytes = self.rows = self.cell_characters = 0
        self.wav_header_bytes = 0
        self.file_rows: dict[str, int] = {}
        self.entries_enumerated = 0
        self.manifest = {
            "schema_version": SCHEMA_VERSION, "scanner_version": SCANNER_VERSION,
            "scan_id": _id(str(root), datetime.now(timezone.utc).isoformat()),
            "scanned_at": datetime.now(timezone.utc).isoformat(), "source_root": str(root),
            "software_versions": {"python": sys.version.split()[0]},
            "policy": {
                "read_only": True, "source_copies": False,
                "parsed_formats": ["csv", "tsv", "xls", "xlsx", "wav_header"],
                "other_formats": "inventory_and_hash_only",
                "source_scope": "only_designated_local_directory", "symlinks": "skip_all",
                "special_files": "skip", "hidden_files": "included", "timezone_default": "unknown",
                "spreadsheet_formulas": "never_execute; XLSX formula text, XLS cached values only",
                "csv_header": "first_record_as_header_hypothesis", "filename_metadata": "unverified_hypotheses",
                "identities": "no_automatic_merge", "hash_scope": "complete_file_or_unknown; never_partial_digest",
                "filesystem_time": "filesystem_observation_not_event_time",
                "wav_header_scope": "bounded_prefix; declared duration only; audio body not validated",
                "wav_header_size_basis": "fstat_of_same_open_descriptor",
                "wav_header_budget": "separate_from_hash_and_table_reads",
                "read_side_effects": "source bytes are not written; OS may update access timestamps",
                "limits": asdict(limits),
            },
            "files": [], "tables": [], "observations": [], "candidate_links": [], "conflicts": [],
            "issues": [], "coverage": {"complete": True, "directories_seen": 0, "entries_seen": 0, "files_seen": 0, "files_hashed": 0, "wav_header_bytes_read": 0, "rows_captured": 0, "cell_characters_captured": 0},
        }
        self._stop = False
        self._reported_limits: set[tuple] = set()

    def issue(self, code: str, path: str | None = None, **details):
        self.manifest["issues"].append({"code": code, "path": path, **details})
        self.manifest["coverage"]["complete"] = False

    def limit(self, code: str, path: str | None = None, **details):
        key = (code, path)
        if key not in self._reported_limits:
            self._reported_limits.add(key)
            self.issue(code, path, **details)

    def observe(self, file: dict, kind: str, raw: Any, *, locator: dict, basis: str, header: Any = None, **extra):
        if len(self.manifest["observations"]) >= self.limits.max_observations:
            self.limit("observation_limit")
            return
        encoded = _json_value(raw)
        obs = {"observation_id": _id(file["file_id"], locator, kind), "file_id": file["file_id"], "kind": kind, "raw_value": encoded,
               "locator": locator, "basis": basis, "status": "unverified"}
        if kind == "phone":
            obs["normalized_candidate"] = _phone(raw)
            if isinstance(raw, (int, float)):
                obs["uncertainty"] = "numeric_spreadsheet_value_may_have_lost_digits_or_leading_zeros"
        elif kind == "timestamp":
            obs.update(_datetime_view(raw))
        elif kind == "duration":
            obs.update(_duration_view(raw, header))
        obs.update(extra)
        self.manifest["observations"].append(obs)
        return obs

    def filename_observations(self, file: dict):
        stem = Path(file["relative_path"]).stem
        timestamp_ranges = []
        pattern = r"(?<!\d)(20\d{2})[-_.]?(\d{2})[-_.]?(\d{2})(?:[T _-]?(\d{2})[:_.-]?(\d{2})(?:[:_.-]?(\d{2}))?)?(?!\d)"
        for match in re.finditer(pattern, stem):
            y, m, d, h, minute, second = match.groups()
            try:
                parsed = datetime(int(y), int(m), int(d), int(h or 0), int(minute or 0), int(second or 0)) if h is not None else date(int(y), int(m), int(d))
            except ValueError:
                continue
            timestamp_ranges.append(match.span())
            self.observe(file, "timestamp", match.group(), locator={"kind": "filename", "start": match.start(), "end": match.end()}, basis="filename_pattern", normalized=parsed.isoformat(), timezone="unknown", interpretation="filename_candidate")
        for match in re.finditer(r"(?<![A-Za-z0-9+])\+?\d(?:[ ()-]?\d){5,14}(?![A-Za-z0-9])", stem):
            if any(match.start() < end and match.end() > start for start, end in timestamp_ranges):
                continue
            self.observe(file, "phone", match.group(), locator={"kind": "filename", "start": match.start(), "end": match.end()}, basis="filename_pattern")

    def walk(self, directory_fd: int, prefix: str = "", depth: int = 0):
        self.manifest["coverage"]["directories_seen"] += 1
        if self.entries_enumerated >= self.limits.max_total_entries:
            self.limit("total_entry_limit", prefix)
            return
        try:
            # Enumerate in lexical order for stable identifiers and reproducible truncation.
            with os.scandir(directory_fd) as entries:
                names = []
                for entry in entries:
                    if self.entries_enumerated >= self.limits.max_total_entries:
                        self.limit("total_entry_limit", prefix)
                        break
                    if len(names) >= self.limits.max_directory_entries:
                        self.issue("directory_entry_limit", prefix, ordering="only_captured_entry_set_sorted; filesystem_enumeration_subset")
                        break
                    names.append(entry.name)
                    self.entries_enumerated += 1
                names.sort()
        except OSError as exc:
            self.issue("directory_unreadable", prefix, error=str(exc))
            return
        for name in names:
            if self._stop:
                return
            rel = f"{prefix}/{name}" if prefix else name
            self.manifest["coverage"]["entries_seen"] += 1
            try:
                st = os.stat(name, dir_fd=directory_fd, follow_symlinks=False)
            except OSError as exc:
                self.issue("entry_unreadable", rel, error=str(exc)); continue
            if stat.S_ISLNK(st.st_mode):
                self.issue("symlink_skipped", rel); continue
            if stat.S_ISDIR(st.st_mode):
                if depth >= self.limits.max_depth:
                    self.issue("depth_limit", rel); continue
                try:
                    child = os.open(name, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW, dir_fd=directory_fd)
                    try:
                        self.walk(child, rel, depth + 1)
                    finally:
                        os.close(child)
                except OSError as exc:
                    self.issue("directory_unreadable", rel, error=str(exc))
            elif stat.S_ISREG(st.st_mode):
                if len(self.manifest["files"]) >= self.limits.max_files:
                    self.issue("file_limit", rel); self._stop = True; return
                self.file(directory_fd, name, rel, st)
            else:
                self.issue("special_file_skipped", rel)

    def file(self, directory_fd: int, name: str, rel: str, st):
        ext = Path(name).suffix.lower()
        file = {"file_id": _id(str(self.root), rel), "relative_path": rel, "name": name, "extension": ext,
                "kind": "table" if ext in {".csv", ".tsv", ".xls", ".xlsx"} else "audio" if ext in _AUDIO else "video" if ext in _VIDEO else "other",
                "size_bytes": st.st_size, "mtime_ns": st.st_mtime_ns, "sha256": None, "hash_status": "unknown", "parse_status": "not_applicable"}
        self.manifest["files"].append(file)
        self.manifest["coverage"]["files_seen"] += 1
        self.filename_observations(file)
        try:
            fd = os.open(name, os.O_RDONLY | os.O_NOFOLLOW | os.O_NONBLOCK, dir_fd=directory_fd)
            try:
                handle = os.fdopen(fd, "rb", buffering=0 if ext == ".wav" else -1)
            except BaseException:
                os.close(fd)
                raise
            with handle:
                before = os.fstat(handle.fileno())
                if (not stat.S_ISREG(before.st_mode) or
                        (before.st_ino, before.st_dev, before.st_size, before.st_mtime_ns, before.st_ctime_ns) !=
                        (st.st_ino, st.st_dev, st.st_size, st.st_mtime_ns, st.st_ctime_ns)):
                    self.issue("source_changed_before_read", rel)
                    self._invalidate_file(file, "unstable", "source_changed_before_read")
                    return
                if st.st_size > self.limits.max_hash_bytes or self.hash_bytes + st.st_size > self.limits.max_total_hash_bytes:
                    self.issue("hash_size_limit", rel, size_bytes=st.st_size)
                else:
                    hasher, consumed = hashlib.sha256(), 0
                    while True:
                        # Limit actual bytes too: a file can grow while being read.
                        allowance = min(self.limits.max_hash_bytes - consumed, self.limits.max_total_hash_bytes - self.hash_bytes)
                        chunk = handle.read(min(1024 * 1024, allowance + 1))
                        if not chunk:
                            file.update(sha256=hasher.hexdigest(), hash_status="complete")
                            self.manifest["coverage"]["files_hashed"] += 1
                            break
                        consumed += len(chunk); self.hash_bytes += len(chunk)
                        if len(chunk) > allowance:
                            self.issue("hash_size_limit_during_read", rel); break
                        hasher.update(chunk)
                handle.seek(0)
                if file["kind"] == "table":
                    if st.st_size > self.limits.max_table_bytes:
                        file["parse_status"] = "skipped_size_limit"; self.issue("table_size_limit", rel)
                    else:
                        raw = handle.read(self.limits.max_table_bytes + 1)
                        if len(raw) > self.limits.max_table_bytes:
                            file["parse_status"] = "skipped_size_limit"; self.issue("table_size_limit_during_read", rel)
                        else:
                            self.parse_table(file, raw)
                elif ext == ".wav":
                    self.wav_header(file, handle, os.fstat(handle.fileno()).st_size)
                after = os.fstat(handle.fileno())
                if (before.st_size, before.st_mtime_ns, before.st_ctime_ns) != (after.st_size, after.st_mtime_ns, after.st_ctime_ns):
                    self._invalidate_file(file, "unstable", "source_changed_during_read")
                    self.issue("source_changed_during_read", rel)
                elif "audio_metadata" in file:
                    file["audio_metadata"]["source_stability"] = "descriptor_stat_unchanged"
        except OSError as exc:
            self._invalidate_file(file, "unreadable", "source_unreadable")
            self.issue("file_unreadable", rel, error=str(exc))

    def _invalidate_file(self, file: dict, status: str, reason: str):
        """Revoke accepted projections without erasing captured raw header declarations."""
        if file["hash_status"] == "complete":
            self.manifest["coverage"]["files_hashed"] -= 1
        file.update(sha256=None, hash_status=status, parse_status=status)
        if "audio_metadata" in file:
            metadata = file["audio_metadata"]
            metadata["source_stability"] = reason
            for key in ("frames", "sample_rate", "channels", "sample_width_bytes", "duration_seconds"):
                metadata.pop(key, None)
            report = metadata["header_probe"]
            if report["status"] == "observed":
                report["status"] = "partial"
            report.update(declared_duration_sec=None, duration_basis=None)
            report["issues"].append({"code": reason, "message": "Source stability check failed; raw declarations retained", "locator": None})
            for observation in self.manifest["observations"]:
                if observation["file_id"] == file["file_id"] and observation["locator"]["kind"] == "wav_header":
                    observation.update(seconds=None, status=reason)

    def wav_header(self, file: dict, handle, descriptor_size: int):
        """Inspect a bounded prefix on the already-owned, unbuffered descriptor.

        Short reads are not EOF. Accounting advances before parsing and remains
        charged if a later read fails. This budget is separate from full hashing.
        """
        rel = file["relative_path"]
        remaining = max(0, self.limits.max_total_wav_header_bytes - self.wav_header_bytes)
        cap = min(self.limits.max_wav_header_bytes, remaining)
        prefix = bytearray()
        eof = False
        read_error = None
        try:
            while len(prefix) < cap:
                requested = min(8192, cap - len(prefix))
                chunk = handle.read(requested)
                if chunk == b"":
                    eof = True
                    break
                if not isinstance(chunk, bytes) or not 0 < len(chunk) <= requested:
                    raise OSError("WAV stream returned invalid read progress")
                self.wav_header_bytes += len(chunk)
                self.manifest["coverage"]["wav_header_bytes_read"] = self.wav_header_bytes
                prefix.extend(chunk)
        except OSError as error:
            read_error = str(error)[:200]
            self.issue("audio_metadata_failed", rel, error=read_error, header_bytes_read=len(prefix))
        if not eof and read_error is None:
            code = "wav_header_total_bytes_limit" if remaining <= self.limits.max_wav_header_bytes else "wav_header_bytes_limit"
            self.limit(code, rel, header_bytes_read=len(prefix), limit_bytes=cap)
        report = probe_wav_header(bytes(prefix), provider_size_bytes=descriptor_size, end_of_input=eof)
        if report["riff_declared_bytes"] is not None and report["riff_declared_bytes"] != descriptor_size:
            if report["status"] not in {"malformed", "unsupported"}:
                report["status"] = "size_mismatch"
            report.update(declared_duration_sec=None, duration_basis=None)
            report["issues"].append({"code": "wav_descriptor_size_mismatch",
                "message": f"RIFF declared size differs from {descriptor_size} bytes observed by fstat", "locator": None})
        if read_error is not None:
            report.update(status="partial", declared_duration_sec=None, duration_basis=None)
        file["audio_metadata"] = {
            "source": "same_descriptor_riff_header", "scope": "header_prefix_only", "body_validated": False,
            "size_basis": "fstat_same_descriptor", "observed_file_size_bytes": descriptor_size,
            "header_bytes_read": len(prefix), "header_probe": report, "source_stability": "not_checked",
        }
        for issue in report["issues"]:
            self.issue(issue["code"], rel, message=issue["message"])
        file["parse_status"] = ("failed" if read_error is not None else
            {"observed": "metadata_only", "partial": "limited", "malformed": "failed",
             "size_mismatch": "failed", "unsupported": "unsupported"}[report["status"]])
        if report["status"] == "observed" and read_error is None:
            seconds = report["declared_duration_sec"]
            file["audio_metadata"].update(frames=report["data_declared_bytes"] // report["block_align_bytes"],
                sample_rate=report["sample_rate_hz"], channels=report["channels"],
                sample_width_bytes=report["bits_per_sample"] // 8, duration_seconds=seconds)
            self.observe(file, "duration", seconds, locator={"kind": "wav_header"}, basis="wav_header",
                header="durationseconds", duration_basis=report["duration_basis"], scope="header_prefix_only",
                body_validated=False, source="same_descriptor_riff_header")

    def add_row(self, file: dict, table: dict, row_index: int, values: list, headers: list | None, *, cells: list | None = None, raw_record: str | None = None, start_line: int | None = None) -> bool:
        if self.file_rows.get(file["file_id"], 0) >= self.limits.max_rows_per_file or self.rows >= self.limits.max_total_rows:
            self.limit("row_limit", file["relative_path"], table=table["name"]); table["complete"] = False; return False
        if self.cell_characters >= self.limits.max_total_cell_characters:
            self.limit("cell_character_limit", file["relative_path"]); table["complete"] = False; return False
        table["rows_seen"] += 1
        captured = []
        if len(values) > self.limits.max_columns:
            self.limit("column_limit", file["relative_path"], row=row_index, total_columns=len(values)); table["complete"] = False
        for col, value in enumerate(values[:self.limits.max_columns]):
            encoded = _json_value(value)
            text = encoded if isinstance(encoded, str) else json.dumps(encoded, ensure_ascii=False)
            remaining = self.limits.max_total_cell_characters - self.cell_characters
            allowed = min(self.limits.max_cell_characters, max(0, remaining))
            truncated = len(text) > allowed
            if truncated:
                encoded = text[:allowed]
                self.limit("cell_value_truncated", file["relative_path"]); table["complete"] = False
            self.cell_characters += min(len(text), allowed)
            cell = {"column": col + 1, "value": encoded, "type": type(value).__name__, "complete": not truncated}
            if cells and col < len(cells):
                cell.update(cells[col])
            captured.append(cell)
            if headers is not None and col < len(headers) and not truncated:
                kind = _kind_for_header(headers[col])
                if kind and value is not None and value != "":
                    self.observe(file, kind, value, locator={"kind": "table_cell", "table_id": table["table_id"], "sheet": table["name"], "row": row_index, "column": col + 1}, basis="column_header_hypothesis", header=headers[col], header_value=_json_value(headers[col]))
        row = {"row_id": _id(table["table_id"], row_index), "row": row_index, "role": "header_hypothesis" if headers is None else "data", "cells": captured}
        if start_line is not None:
            row["start_line"] = start_line
        if raw_record is not None:
            # Raw CSV record is useful for quoting and malformed-row diagnosis. It
            # shares the same character budget; duplicated values are not unlimited.
            allowed = min(self.limits.max_cell_characters * self.limits.max_columns, max(0, self.limits.max_total_cell_characters - self.cell_characters))
            row["raw_record"] = raw_record[:allowed]
            row["raw_record_complete"] = len(raw_record) <= allowed
            self.cell_characters += min(len(raw_record), allowed)
            if len(raw_record) > allowed:
                self.limit("raw_record_truncated", file["relative_path"]); table["complete"] = False
        if headers is not None and len(values) != len(headers):
            self.issue("row_width_mismatch", file["relative_path"], sheet=table["name"], row=row_index, expected=len(headers), actual=len(values))
            row["issue"] = "row_width_mismatch"
        table["rows"].append(row); self.rows += 1
        self.file_rows[file["file_id"]] = self.file_rows.get(file["file_id"], 0) + 1
        return True

    def table(self, file: dict, name: str, **extra) -> dict:
        table = {"table_id": _id(file["file_id"], name), "file_id": file["file_id"], "name": name, "rows_seen": 0, "complete": True, "rows": [], **extra}
        self.manifest["tables"].append(table)
        return table

    def parse_table(self, file: dict, raw: bytes):
        file["parse_status"] = "complete"
        try:
            if file["extension"] in {".csv", ".tsv"}:
                self.csv(file, raw)
            elif file["extension"] == ".xlsx":
                self.xlsx(file, raw)
            else:
                self.xls(file, raw)
        except ImportError as exc:
            file["parse_status"] = "dependency_unavailable"; self.issue("parser_dependency_unavailable", file["relative_path"], dependency=exc.name)
        except Exception as exc:
            file["parse_status"] = "failed"; self.issue("table_parse_failed", file["relative_path"], error=f"{type(exc).__name__}: {exc}")
        if any(not table["complete"] for table in self.manifest["tables"] if table["file_id"] == file["file_id"]) and file["parse_status"] == "complete":
            file["parse_status"] = "limited"

    def csv(self, file: dict, raw: bytes):
        encoding = "utf-16" if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
        try:
            text = raw.decode(encoding)
        except UnicodeDecodeError:
            encoding = "cp1250"
            text = raw.decode(encoding, errors="strict")
            self.issue("csv_encoding_inferred", file["relative_path"], encoding=encoding, uncertainty="fallback_not_verified")
        lines = text.splitlines(keepends=True)
        delimiter, start = "\t" if file["extension"] == ".tsv" else None, 0
        if lines and re.fullmatch(r"sep=([^\r\n])\r?\n?", lines[0], re.IGNORECASE):
            delimiter = lines[0][4]; start = 1
        if delimiter is None:
            try:
                delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
            except csv.Error:
                first_line = lines[start] if len(lines) > start else ""
                # A malformed later record must not erase a clear header dialect.
                counts = {candidate: first_line.count(candidate) for candidate in (",", ";", "\t", "|")}
                delimiter = max(counts, key=counts.get) if any(counts.values()) else ","
        table = self.table(file, "csv", encoding=encoding, delimiter=delimiter, preamble=lines[0] if start else None, parser="python.csv", parser_version=sys.version.split()[0])
        reader = csv.reader(io.StringIO("".join(lines[start:])), delimiter=delimiter, strict=True)
        headers, previous_line = None, 0
        try:
            for index, values in enumerate(reader, start=1):
                raw_record = "".join(lines[start + previous_line:start + reader.line_num])
                start_line = start + previous_line + 1
                previous_line = reader.line_num
                if not self.add_row(file, table, index, values, headers, raw_record=raw_record, start_line=start_line):
                    break
                if headers is None:
                    headers = values

        except csv.Error as exc:
            fragment = "".join(lines[start + previous_line:])
            remaining = max(0, self.limits.max_total_cell_characters - self.cell_characters)
            allowed = min(self.limits.max_cell_characters, remaining)
            self.cell_characters += min(len(fragment), allowed)
            table["complete"] = False
            file["parse_status"] = "failed"
            self.issue("csv_record_parse_failed", file["relative_path"], start_line=start + previous_line + 1, error=str(exc), raw_fragment=fragment[:allowed], raw_fragment_complete=len(fragment) <= allowed)
    def xlsx(self, file: dict, raw: bytes):
        import openpyxl
        self.manifest["software_versions"]["openpyxl"] = openpyxl.__version__
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            expanded = sum(info.file_size for info in archive.infolist())
            if expanded > self.limits.max_xlsx_expanded_bytes:
                file["parse_status"] = "skipped_expansion_limit"; self.issue("xlsx_expansion_limit", file["relative_path"], expanded_bytes=expanded); return
        workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            for sheet_index, sheet in enumerate(workbook):
                if sheet_index >= self.limits.max_sheets_per_file:
                    self.issue("sheet_limit", file["relative_path"], next_sheet=sheet.title)
                    file["parse_status"] = "limited"
                    break
                table = self.table(file, sheet.title, formula_visibility="formula_text; not_evaluated; cached_values_not_exposed")
                if sheet.max_column and sheet.max_column > self.limits.max_columns:
                    self.limit("column_range_not_scanned", file["relative_path"], sheet=sheet.title, declared_columns=sheet.max_column)
                    table["complete"] = False
                headers = None
                for index, row in enumerate(sheet.iter_rows(max_col=self.limits.max_columns + 1), start=1):
                    values = [cell.value for cell in row]
                    while values and values[-1] is None:
                        values.pop()
                    cells = [{"cell_type": cell.data_type, "number_format": cell.number_format, "coordinate": getattr(cell, "coordinate", None)} for cell in row[:len(values)]]
                    if not self.add_row(file, table, index, values, headers, cells=cells):
                        break
                    if headers is None:
                        headers = values
        finally:
            workbook.close()

    def xls(self, file: dict, raw: bytes):
        import xlrd
        self.manifest["software_versions"]["xlrd"] = xlrd.__version__
        workbook = xlrd.open_workbook(file_contents=raw, on_demand=True, ragged_rows=True)
        try:
            for sheet_index in range(workbook.nsheets):
                if sheet_index >= self.limits.max_sheets_per_file:
                    self.issue("sheet_limit", file["relative_path"], next_sheet_index=sheet_index)
                    file["parse_status"] = "limited"
                    break
                sheet = workbook.sheet_by_index(sheet_index)
                table = self.table(file, sheet.name, formula_visibility="cached_values_only; formula_text_unavailable", workbook_date_mode=workbook.datemode)
                headers = None
                for index in range(sheet.nrows):
                    values = sheet.row_values(index)
                    cells = [{"cell_type": sheet.cell_type(index, col)} for col in range(len(values))]
                    # Keep the raw Excel serial number; expose a separate date candidate.
                    for col, kind in enumerate(sheet.row_types(index)):
                        if kind == xlrd.XL_CELL_DATE:
                            try:
                                cells[col]["date_serial_candidate"] = xlrd.xldate_as_datetime(values[col], workbook.datemode).isoformat()
                                cells[col]["timezone"] = "unknown"
                            except (ValueError, OverflowError) as exc:
                                cells[col]["date_interpretation_error"] = str(exc)
                                self.issue("xls_date_interpretation_failed", file["relative_path"], sheet=sheet.name, row=index + 1, column=col + 1)
                    if not self.add_row(file, table, index + 1, values, headers, cells=cells):
                        break
                    if headers is None:
                        headers = values
        finally:
            workbook.release_resources()

    def correlate(self):
        files = self.manifest["files"]
        by_name: dict[str, list[dict]] = {}
        for file in files:
            by_name.setdefault(file["name"], []).append(file)
        for name, group in by_name.items():
            if len(group) > 1:
                self.manifest["conflicts"].append({"kind": "filename_collision", "name": name, "file_ids": [f["file_id"] for f in group], "resolution": "unresolved; distinct_paths_retained"})
        observations = self.manifest["observations"]
        row_phones: dict[tuple, list[dict]] = {}
        filename_phones: dict[str, list[dict]] = {}
        for obs in observations:
            loc = obs["locator"]
            if obs["kind"] == "phone":
                if loc["kind"] == "table_cell":
                    row_phones.setdefault((loc["table_id"], loc["row"]), []).append(obs)
                elif loc["kind"] == "filename":
                    filename_phones.setdefault(obs["file_id"], []).append(obs)
        for ref in observations:
            if ref["kind"] != "file_reference" or not isinstance(ref["raw_value"], str):
                continue
            loc = ref["locator"]
            value = ref["raw_value"].replace("\\", "/")
            candidates = by_name.get(value.rsplit("/", 1)[-1], [])
            for candidate in candidates:
                if len(self.manifest["candidate_links"]) >= self.limits.max_candidate_links:
                    self.limit("candidate_link_limit"); return
                self.manifest["candidate_links"].append({"observation_id": ref["observation_id"], "target_file_id": candidate["file_id"], "basis": "exact_basename_reference", "status": "candidate", "ambiguous": len(candidates) > 1})
                row_values = row_phones.get((loc["table_id"], loc["row"]), [])
                name_values = filename_phones.get(candidate["file_id"], [])
                row_numbers = sorted({o["normalized_candidate"] for o in row_values if o["normalized_candidate"]})
                name_numbers = sorted({o["normalized_candidate"] for o in name_values if o["normalized_candidate"]})
                if row_numbers and name_numbers and set(row_numbers).isdisjoint(name_numbers):
                    self.manifest["conflicts"].append({"kind": "phone_candidates_disagree", "file_id": candidate["file_id"], "reference_observation_id": ref["observation_id"], "table_phone_observation_ids": [o["observation_id"] for o in row_values], "filename_phone_observation_ids": [o["observation_id"] for o in name_values], "table_candidates": row_numbers, "filename_candidates": name_numbers, "resolution": "unresolved; neither_source_selected"})


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, float) and (value != value or abs(value) == float("inf")):
        return {"representation": repr(value), "non_finite": True}
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return {"representation": repr(value), "type": type(value).__name__}


def scan_sources(root: str | Path, *, limits: ScanLimits | None = None) -> dict:
    """Inventory a designated local directory without copying or writing sources.

    IDs identify source root + exact relative path, not basename or presumed person.
    Hashes identify bytes independently. Repeat scans produce stable file, table and
    observation IDs; scan_id/scanned_at identify the new acquisition observation.
    """
    requested = Path(root).expanduser()
    if requested.is_symlink():
        raise ValueError("Source root must be a real directory, not a symlink")
    source = requested.resolve(strict=True)
    if not source.is_dir():
        raise ValueError("Source root must be a directory")
    if not all(hasattr(os, name) for name in ("O_NOFOLLOW", "O_DIRECTORY")):
        raise RuntimeError("Scanner requires safe no-follow directory/file descriptors")
    scan = _Scan(source, limits or ScanLimits())
    root_fd = os.open(source, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW)
    try:
        scan.walk(root_fd)
    finally:
        os.close(root_fd)
    scan.correlate()
    scan.manifest["coverage"].update(rows_captured=scan.rows, cell_characters_captured=scan.cell_characters, hash_bytes_read=scan.hash_bytes, wav_header_bytes_read=scan.wav_header_bytes,
                                       entries_enumerated=scan.entries_enumerated, inventory_files=len(scan.manifest["files"]), tables=len(scan.manifest["tables"]), observations=len(scan.manifest["observations"]))
    return scan.manifest


def export_manifest(manifest: dict, output_path: str | Path, *, source_root: str | Path | None = None, max_bytes: int = 32 * 1024 * 1024) -> Path:
    """Atomically write one small metadata manifest outside the scanned tree."""
    sources = {Path(manifest["source_root"]).resolve(strict=False)}
    if source_root is not None:
        sources.add(Path(source_root).resolve(strict=False))
    output = Path(output_path).expanduser()
    if output.is_symlink():
        raise ValueError("Output cannot be a symlink")
    output = output.resolve(strict=False)
    if any(output == source or source in output.parents for source in sources):
        raise ValueError("Manifest output must be outside the scanned source root")
    if max_bytes < 1:
        raise ValueError("max_bytes must be positive")
    payload = json.dumps(manifest, ensure_ascii=False, indent=2, allow_nan=False).encode("utf-8") + b"\n"
    if len(payload) > max_bytes:
        raise ValueError(f"Manifest exceeds output size limit ({len(payload)} > {max_bytes}); reduce scan limits")
    output.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix=".ageds-scan-", suffix=".tmp", dir=output.parent)
    try:
        with os.fdopen(fd, "wb") as handle:
            handle.write(payload); handle.flush(); os.fsync(handle.fileno())
        os.replace(temporary, output)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
    return output
