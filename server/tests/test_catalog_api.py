"""HTTP API and server-rendered /catalog page on the synthetic corpus."""
from __future__ import annotations

import csv
import io
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from server.app import main
from server.app.catalog import describe
from server.tests.test_catalog_scan import make_world, tree_state


@pytest.fixture
def api(tmp_path, monkeypatch):
    world = make_world(tmp_path, monkeypatch)
    monkeypatch.setattr(main, 'settings', world.settings)
    with TestClient(main.app) as client:
        world.client = client
        yield world


def test_page_renders_tree_filters_and_escapes_source_names(api):
    response = api.client.get('/catalog')
    assert response.status_code == 200
    page = response.text
    assert '<a href="/catalog">Katalog</a>' in page
    assert 'Phone A (synthetic)' in page and 'Drive B (synthetic)' in page
    assert '&lt;script&gt;alert(1).txt' in page
    assert '<script>alert' not in page
    # Only the page's own module plus the shared shell (nav) scripts from /static.
    shell = ('<script type="application/json" id="ageds-shell-config">', '<script type="module" src="/static/workbench-shell.mjs">')
    assert page.count('<script') == 1 + sum(page.count(tag) for tag in shell)
    assert page.count('<script type="module" src="/static/catalog.mjs">') == 1
    selected = api.client.get('/catalog', params={'collection': api.a['id'], 'conflict': 'true'}).text
    for locator in api.expected['type_conflicts']:
        assert locator in selected
    assert 'Pokrycie i luki' in selected and '2024-01-09 → 2024-09-01' in selected
    assert 'href="/api/catalog/export?format=xlsx&amp;collectionId=' in selected
    paged = api.client.get('/catalog', params={'collection': api.a['id'], 'limit': 5}).text
    assert 'rel="next"' in paged
    assert 'Katalog' in api.client.get('/').text


def test_collection_crud_and_audit(api):
    client = api.client
    created = client.post('/api/catalog/collections', json={'label': 'Laptop export', 'priority': 1}).json()
    patched = client.patch(f'/api/catalog/collections/{created["id"]}', json={'notes': 'n1', 'priority': None}).json()
    assert patched['notes'] == 'n1' and patched['priority'] is None
    assert patched['history'][0]['details'] == {'previous': {'notes': None, 'priority': 1}, 'new': {'notes': 'n1', 'priority': None}}
    assert client.patch(f'/api/catalog/collections/{created["id"]}', json={'case_id': 3}).status_code == 422
    assert client.patch(f'/api/catalog/collections/{api.a["id"]}', json={'parent_collection_id': api.a['id']}).status_code == 422
    assert client.get('/api/catalog/collections/999999').status_code == 404
    assert client.post('/api/catalog/collections', json={'label': ''}).status_code == 422
    labels = [c['label'] for c in client.get('/api/catalog/collections').json()['collections']]
    assert 'Laptop export' in labels


def test_http_scan_is_confined_to_configured_roots(api):
    client = api.client
    body = {'collectionId': api.b['id'], 'rootId': 0, 'relativePath': 'drive_b'}
    result = client.post('/api/catalog/scans', json=body)
    assert result.status_code == 200 and result.json()['status'] == 'completed'
    assert client.post('/api/catalog/scans', json={**body, 'relativePath': '../'}).status_code == 422
    assert client.post('/api/catalog/scans', json={**body, 'relativePath': 'phone_a/link-to-docs'}).status_code == 422
    assert client.post('/api/catalog/scans', json={**body, 'rootId': 5}).status_code == 404
    assert client.post('/api/catalog/scans', json={**body, 'limits': {'max_bogus': 1}}).status_code == 422
    assert client.post('/api/catalog/scans', json={**body, 'limits': {'max_files': 0}}).status_code == 422
    scans = client.get('/api/catalog/scans', params={'collectionId': api.b['id']}).json()['scans']
    assert [s['current'] for s in scans] == [True, False]
    assert tree_state(api.corpus) == api.before


def test_entries_groups_collisions_coverage(api):
    client = api.client
    first = client.get('/api/catalog/entries', params={'collectionId': api.a['id'], 'limit': 7}).json()
    assert len(first['entries']) == 7 and first['has_more']
    second = client.get('/api/catalog/entries', params={'collectionId': api.a['id'], 'limit': 7, 'afterId': first['next_after_id']}).json()
    assert second['entries'][0]['id'] > first['entries'][-1]['id']
    assert client.get('/api/catalog/entries', params={'limit': 501}).status_code == 422
    assert client.get('/api/catalog/entries', params={'status': 'nope'}).status_code == 422
    encrypted = client.get('/api/catalog/entries', params={'status': 'encrypted'}).json()['entries']
    assert [e['locator'] for e in encrypted] == [api.expected['encrypted_member']]
    detail = client.get(f'/api/catalog/entries/{encrypted[0]["id"]}').json()
    assert detail['containers'][0]['locator'] == 'archives/locked.zip'
    groups = client.get('/api/catalog/groups').json()['groups']
    group = client.get(f'/api/catalog/groups/{groups[0]["id"]}').json()
    assert len(group['catalog_entries']) == 3 and 'not authorship' in group['note']
    names = {c['name'] for c in client.get('/api/catalog/collisions').json()['collisions']}
    assert {'todo.txt', 'same-name.txt'} <= names
    report = client.get('/api/catalog/coverage', params={'collectionId': api.a['id'], 'minGapDays': 100}).json()
    assert report['collections'][0]['totals']['gaps'] == 1
    assert client.get('/api/catalog/coverage', params={'minGapDays': 0}).status_code == 422


def test_export_downloads_are_openable_and_inert(api):
    client = api.client
    response = client.get('/api/catalog/export', params={'format': 'csv'})
    assert response.status_code == 200 and response.headers['cache-control'] == 'no-store'
    rows = list(csv.DictReader(io.StringIO(response.text)))
    assert len(rows) == int(response.headers['x-ageds-rows']) and response.headers['x-ageds-truncated'] == 'false'
    assert any(r['locator'] == 'archives/evil.zip!/tab%09name.txt' and r['name'] == 'tab\\x09name.txt' for r in rows)
    workbook = load_workbook(io.BytesIO(client.get('/api/catalog/export', params={'format': 'xlsx'}).content))
    assert {'entries', 'coverage', 'duplicates', 'collisions', 'about'} <= set(workbook.sheetnames)
    html = client.get('/api/catalog/export', params={'format': 'html', 'collectionId': api.a['id']}).text
    assert '&lt;script&gt;alert(1).txt' in html and '<script' not in html
    truncated = client.get('/api/catalog/export', params={'format': 'csv', 'maxRows': 3})
    assert len(list(csv.DictReader(io.StringIO(truncated.text)))) == 3 and truncated.headers['x-ageds-truncated'] == 'true'


def test_marks_explain_decisions_and_ingest(api):
    client = api.client
    assert client.post('/api/catalog/marks', json={'collectionId': api.a['id'], 'locator': 'docs/report.pdf',
                                                   'priority': 4, 'note': 'check'}).status_code == 200
    assert client.post('/api/catalog/marks', json={'collectionId': api.a['id'], 'locator': 'nope', 'priority': 1}).status_code == 404
    none = client.post(f'/api/catalog/collections/{api.a["id"]}/explain', json={}).json()
    assert none['status'] == 'not_run' and 'profile' in none
    run = client.post(f'/api/catalog/collections/{api.a["id"]}/explain', json={'provider': 'fake'}).json()
    sidecar = api.settings.data_dir / 'catalog-annotations' / f'collection-{api.a["id"]}.annotations.ndjson'
    assert run['sidecar'] == str(sidecar) and sidecar.exists()
    annotation = run['annotation']['annotation_id']
    assert client.post(f'/api/catalog/annotations/{annotation}/decision', json={'decision': 'promoted'}).status_code == 422
    assert client.post(f'/api/catalog/annotations/{annotation}/decision', json={'decision': 'promoted', 'actor': 'reviewer'}).status_code == 200
    listed = client.get('/api/catalog/annotations', params={'collectionId': api.a['id']}).json()['annotations']
    assert listed[0]['current_status'] == 'promoted'
    assert client.post(f'/api/catalog/collections/{api.a["id"]}/explain', json={'provider': 'unknown'}).status_code == 422
    entry = client.get('/api/catalog/entries', params={'q': 'deep.txt'}).json()['entries'][0]
    result = client.post('/api/catalog/ingest', json={'entryIds': [entry['id']]}).json()
    assert result['ingested'] == 1
    artifact = result['results'][0]['artifact_id']
    observations = client.get(f'/api/artifacts/{artifact}/source-observations').json()
    assert observations and observations[0]['source_locator'].endswith(api.expected['zip_depth3'])
    assert client.get(f'/api/catalog/entries/{entry["id"]}').json()['ingested'][0]['artifact_id'] == artifact
    assert tree_state(api.corpus) == api.before


def test_external_provider_requires_explicit_opt_in(api, monkeypatch):
    class External(describe.FakeProvider):
        name = 'external-test'
        external = True
    monkeypatch.setitem(describe.PROVIDERS, 'external-test', External())
    url = f'/api/catalog/collections/{api.a["id"]}/explain'
    assert api.client.post(url, json={'provider': 'external-test'}).status_code == 403
    allowed = api.client.post(url, json={'provider': 'external-test', 'allowExternal': True}).json()
    assert allowed['annotation']['external'] is True and allowed['annotation']['opt_in']['allow_external'] is True
