from __future__ import annotations

import hashlib
import json
import os
import struct
import tempfile
import unittest
import wave
from pathlib import Path

from server.app.scanner import ScanLimits, export_manifest, scan_sources


def digest(path: Path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def simple_xls(path: Path):
    """Tiny BIFF8 fixture made directly, independent of the parser under test."""
    record = lambda identifier, data: struct.pack('<HH', identifier, len(data)) + data
    bof = lambda kind: record(0x809, struct.pack('<HHHHII', 0x600, kind, 0xdbb, 0x7cc, 0x41, 6))
    end = record(0xa, b'')
    name = b'Calls'
    bound = lambda position: record(0x85, struct.pack('<IBBBB', position, 0, 0, len(name), 0) + name)
    globals_ = bof(5) + bound(0) + end
    sheet = bof(0x10) + record(0x200, struct.pack('<IIHHH', 0, 2, 0, 2, 0))
    for row, values in enumerate([['Phone', 'Date'], ['+48123456789', '2026-09-30 08:15']]):
        for col, value in enumerate(values):
            encoded = value.encode('latin1')
            sheet += record(0x204, struct.pack('<HHHHB', row, col, 0, len(encoded), 0) + encoded)
    path.write_bytes(bof(5) + bound(len(globals_)) + end + sheet + end)


class ScannerTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.base = Path(self.tmp.name)
        self.source = self.base / 'source'
        self.source.mkdir()

    def tearDown(self):
        self.tmp.cleanup()

    def test_read_only_repeat_scan_path_identity_and_collision(self):
        for folder, data in [('one', b'first'), ('two', b'second')]:
            directory = self.source / folder
            directory.mkdir()
            (directory / 'call_+48123456789.mp3').write_bytes(data)
        paths = sorted(self.source.rglob('*.mp3'))
        before = {str(p): (digest(p), p.stat().st_size, p.stat().st_mtime_ns) for p in paths}
        first, second = scan_sources(self.source), scan_sources(self.source)
        after = {str(p): (digest(p), p.stat().st_size, p.stat().st_mtime_ns) for p in paths}
        self.assertEqual(before, after)
        self.assertEqual([f['file_id'] for f in first['files']], [f['file_id'] for f in second['files']])
        self.assertEqual(len({f['file_id'] for f in first['files']}), 2)
        self.assertEqual(len({f['sha256'] for f in first['files']}), 2)
        self.assertNotEqual(first['scan_id'], second['scan_id'])
        self.assertTrue(any(c['kind'] == 'filename_collision' for c in first['conflicts']))
        self.assertEqual(len(list(self.source.rglob('*'))), 4)

    def test_csv_raw_unicode_bad_row_and_conflicting_phone_hypotheses(self):
        name = 'call_+48123456789_20260930_081530.mp3'
        (self.source / name).write_bytes(b'not actual mp3, inventory only')
        text = f'Numer telefonu;Data;DurationSeconds;Plik;Nazwa\n+48987654321;2026-09-30 08:15:30;30,5;{name};Żółć\n+48000000000;bad-date\n'
        (self.source / 'billing.csv').write_text(text, encoding='utf-8')
        result = scan_sources(self.source)
        table = result['tables'][0]
        self.assertEqual(table['delimiter'], ';')
        self.assertEqual(table['rows'][1]['cells'][4]['value'], 'Żółć')
        self.assertIn('30,5', table['rows'][1]['raw_record'])
        self.assertEqual(table['rows'][2]['issue'], 'row_width_mismatch')
        dates = [o for o in result['observations'] if o['kind'] == 'timestamp']
        self.assertTrue(all(o['timezone'] == 'unknown' for o in dates))
        self.assertTrue(any(o['normalized'] == '2026-09-30T08:15:30' for o in dates))
        self.assertTrue(any(o['normalized'] is None and o['raw_value'] == 'bad-date' for o in dates))
        durations = [o for o in result['observations'] if o['kind'] == 'duration']
        self.assertEqual(durations[0]['seconds'], 30.5)
        self.assertTrue(any(c['kind'] == 'phone_candidates_disagree' for c in result['conflicts']))
        self.assertEqual(result['candidate_links'][0]['status'], 'candidate')
        self.assertFalse(result['coverage']['complete'])

    def test_csv_cp1250_is_explicit_inference_and_multiline_raw_record_preserved(self):
        content = 'Name;Phone\n"Żółć\nŁódź";+48123456789\n'
        (self.source / 'polish.csv').write_bytes(content.encode('cp1250'))
        result = scan_sources(self.source)
        table = result['tables'][0]
        self.assertEqual(table['encoding'], 'cp1250')
        self.assertEqual(table['rows'][1]['cells'][0]['value'], 'Żółć\nŁódź')
        self.assertIn('"Żółć\nŁódź"', table['rows'][1]['raw_record'])
        self.assertTrue(any(i['code'] == 'csv_encoding_inferred' for i in result['issues']))

    def test_csv_utf16_excel_separator(self):
        (self.source / 'csv.csv').write_bytes('sep=;\r\nPhone;Date\r\n+48123456789;2026-09-30T08:15:00+02:00\r\n'.encode('utf-16'))
        result = scan_sources(self.source)
        self.assertEqual(result['tables'][0]['preamble'], 'sep=;\r\n')
        self.assertEqual(result['tables'][0]['rows'][0]['start_line'], 2)
        timestamp = next(o for o in result['observations'] if o['kind'] == 'timestamp')
        self.assertNotEqual(timestamp['timezone'], 'unknown')

    def test_xlsx_keeps_formula_text_numeric_phone_uncertainty_and_separate_dates(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest('optional openpyxl unavailable')
        from datetime import datetime
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = 'Billing'
        sheet.append(['Phone', 'Date', 'formula'])
        sheet.append([48123456789, datetime(2026, 9, 30, 8, 15), '=1+2'])
        workbook.save(self.source / 'billing.xlsx')
        before = digest(self.source / 'billing.xlsx')
        result = scan_sources(self.source)
        self.assertEqual(digest(self.source / 'billing.xlsx'), before)
        table = result['tables'][0]
        self.assertIn('cached_values_not_exposed', table['formula_visibility'])
        self.assertEqual(table['rows'][1]['cells'][2]['value'], '=1+2')
        self.assertEqual(table['rows'][1]['cells'][2]['cell_type'], 'f')
        phone = next(o for o in result['observations'] if o['kind'] == 'phone')
        self.assertIsNone(phone['normalized_candidate'])
        self.assertIn('lost_digits', phone['uncertainty'])
        timestamp = next(o for o in result['observations'] if o['kind'] == 'timestamp')
        self.assertEqual(timestamp['timezone'], 'unknown')

    def test_xls_real_biff_cached_value_limitation_explicit(self):
        try:
            import xlrd
        except ImportError:
            self.skipTest('optional xlrd unavailable')
        simple_xls(self.source / 'billing.xls')
        result = scan_sources(self.source)
        table = result['tables'][0]
        self.assertEqual(table['rows'][1]['cells'][0]['value'], '+48123456789')
        self.assertIn('formula_text_unavailable', table['formula_visibility'])
        self.assertEqual(next(o for o in result['observations'] if o['kind'] == 'phone')['normalized_candidate'], '+48123456789')

    def test_symlinks_cycles_special_files_never_read(self):
        outside = self.base / 'outside'
        outside.mkdir()
        (outside / 'secret.csv').write_text('Phone\n+48111111111\n')
        (self.source / 'outside-link').symlink_to(outside, target_is_directory=True)
        (self.source / 'cycle').symlink_to(self.source, target_is_directory=True)
        (self.source / 'file-link.csv').symlink_to(outside / 'secret.csv')
        os.mkfifo(self.source / 'pipe')
        result = scan_sources(self.source)
        self.assertEqual(result['files'], [])
        self.assertEqual(sum(i['code'] == 'symlink_skipped' for i in result['issues']), 3)
        self.assertTrue(any(i['code'] == 'special_file_skipped' for i in result['issues']))
        with self.assertRaises(ValueError):
            scan_sources(self.source / 'outside-link')

    def test_bounds_capture_and_hash_limits_explicit(self):
        (self.source / 'first.csv').write_text('Phone,Name\n+48123456789,' + 'X' * 30 + '\n+48222222222,second\n')
        result = scan_sources(self.source, limits=ScanLimits(max_hash_bytes=3, max_rows_per_file=2, max_cell_characters=8))
        file = result['files'][0]
        self.assertIsNone(file['sha256'])
        self.assertEqual(len(result['tables'][0]['rows']), 2)
        self.assertEqual(file['parse_status'], 'limited')
        codes = {i['code'] for i in result['issues']}
        self.assertTrue({'hash_size_limit', 'row_limit', 'cell_value_truncated'} <= codes)
        json.dumps(result, allow_nan=False)

    def test_total_rows_limit_not_reset_between_xlsx_sheets(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest('optional openpyxl unavailable')
        workbook = openpyxl.Workbook()
        for sheet in [workbook.active, workbook.create_sheet('other')]:
            sheet.append(['Phone']); sheet.append(['+48123456789'])
        workbook.save(self.source / 'many.xlsx')
        result = scan_sources(self.source, limits=ScanLimits(max_rows_per_file=3))
        self.assertEqual(result['coverage']['rows_captured'], 3)
        self.assertTrue(any(i['code'] == 'row_limit' for i in result['issues']))

    def test_corrupt_spreadsheets_report_failure_and_keep_inventory_hash(self):
        (self.source / 'bad.xlsx').write_bytes(b'broken data')
        result = scan_sources(self.source)
        self.assertEqual(result['files'][0]['parse_status'], 'failed')
        self.assertEqual(result['files'][0]['hash_status'], 'complete')
        self.assertTrue(any(i['code'] == 'table_parse_failed' for i in result['issues']))

    def test_wav_duration_from_header_separate_from_filename(self):
        with wave.open(str(self.source / 'clip.wav'), 'wb') as audio:
            audio.setnchannels(1); audio.setsampwidth(2); audio.setframerate(8000)
            audio.writeframes(b'\0\0' * 4000)
        result = scan_sources(self.source)
        file = result['files'][0]
        self.assertEqual(file['audio_metadata']['duration_seconds'], .5)
        duration = next(o for o in result['observations'] if o['kind'] == 'duration')
        self.assertEqual(duration['basis'], 'wav_header')
        self.assertEqual(duration['seconds'], .5)

    def test_output_must_be_outside_sources_and_size_limit_is_transactional(self):
        (self.source / 'sample.csv').write_text('Phone\n+48123456789\n')
        result = scan_sources(self.source)
        with self.assertRaises(ValueError):
            export_manifest(result, self.source / 'manifest.json')
        (self.base / 'link').symlink_to(self.source, target_is_directory=True)
        with self.assertRaises(ValueError):
            export_manifest(result, self.base / 'link' / 'manifest.json')
        output = self.base / 'manifest.json'
        self.assertEqual(export_manifest(result, output), output)
        previous = output.read_bytes()
        with self.assertRaises(ValueError):
            export_manifest(result, output, max_bytes=3)
        self.assertEqual(output.read_bytes(), previous)
        self.assertEqual(json.loads(previous)['schema_version'], 'ageds.source-scan/v1')

    def test_file_limit_directory_depth_and_observation_limit(self):
        (self.source / 'call_+48123456789_20260930_081530.mp3').write_bytes(b'a')
        (self.source / 'second.txt').write_bytes(b'b')
        result = scan_sources(self.source, limits=ScanLimits(max_files=1, max_observations=1))
        self.assertEqual(len(result['files']), 1)
        self.assertEqual(len(result['observations']), 1)
        self.assertIn('observation_limit', {i['code'] for i in result['issues']})
        self.assertIn('file_limit', {i['code'] for i in result['issues']})
        deep = self.source / 'one' / 'two'
        deep.mkdir(parents=True)
        (deep / 'skip.txt').write_text('deep')
        result = scan_sources(self.source, limits=ScanLimits(max_depth=1))
        self.assertIn('depth_limit', {i['code'] for i in result['issues']})

    def test_malformed_csv_captures_failed_record_fragment_without_rewriting(self):
        path = self.source / 'bad.csv'
        raw = 'Phone,Name\n+48123456789,"broken\n'
        path.write_text(raw)
        before = digest(path)
        result = scan_sources(self.source)
        self.assertEqual(digest(path), before)
        issue = next(i for i in result['issues'] if i['code'] == 'csv_record_parse_failed')
        self.assertEqual(issue['start_line'], 2)
        self.assertEqual(issue['raw_fragment'], '+48123456789,"broken\n')
        self.assertEqual(result['files'][0]['parse_status'], 'failed')

    def test_export_override_cannot_bypass_actual_scanned_source(self):
        result = scan_sources(self.source)
        with self.assertRaises(ValueError):
            export_manifest(result, self.source / 'evil.json', source_root=self.base / 'nonexistent')
        outside = self.base / 'outside'
        outside.mkdir()
        with self.assertRaises(ValueError):
            export_manifest(result, self.source / 'evil.json', source_root=outside)

    def test_observation_cap_never_mutates_previous_timestamp(self):
        (self.source / '2026-09-29.wav').write_bytes(b'broken')
        (self.source / '2026-09-30.wav').write_bytes(b'broken')
        result = scan_sources(self.source, limits=ScanLimits(max_observations=1))
        self.assertEqual(result['observations'][0]['raw_value'], '2026-09-29')
        self.assertEqual(result['observations'][0]['normalized'], '2026-09-29')

    def test_xlsx_expansion_and_sheet_limits_are_reported(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest('optional openpyxl unavailable')
        workbook = openpyxl.Workbook()
        workbook.active.append(['Phone'])
        workbook.create_sheet('second').append(['Phone'])
        workbook.save(self.source / 'many.xlsx')
        small = scan_sources(self.source, limits=ScanLimits(max_xlsx_expanded_bytes=1))
        self.assertEqual(small['files'][0]['parse_status'], 'skipped_expansion_limit')
        limited = scan_sources(self.source, limits=ScanLimits(max_sheets_per_file=1))
        self.assertEqual(len(limited['tables']), 1)
        self.assertEqual(limited['files'][0]['parse_status'], 'limited')
        self.assertIn('sheet_limit', {i['code'] for i in limited['issues']})

    def test_basename_reference_collision_retains_both_candidates(self):
        for dirname in ['left', 'right']:
            directory = self.source / dirname
            directory.mkdir()
            (directory / 'clip.mp3').write_bytes(dirname.encode())
        (self.source / 'billing.csv').write_text('File,Phone\nclip.mp3,+48123456789\n')
        result = scan_sources(self.source)
        self.assertEqual(len(result['candidate_links']), 2)
        self.assertTrue(all(c['ambiguous'] for c in result['candidate_links']))
        self.assertEqual(len({c['target_file_id'] for c in result['candidate_links']}), 2)

    def test_directory_enumeration_bound_is_explicit(self):
        for index in range(4):
            (self.source / f'{index}.txt').write_text(str(index))
        result = scan_sources(self.source, limits=ScanLimits(max_directory_entries=2))
        self.assertEqual(len(result['files']), 2)
        self.assertIn('directory_entry_limit', {i['code'] for i in result['issues']})

    def test_total_entry_limit_bounds_tree_of_empty_directories(self):
        for parent in ['a', 'b']:
            for child in ['one', 'two']:
                (self.source / parent / child).mkdir(parents=True)
        result = scan_sources(self.source, limits=ScanLimits(max_total_entries=3))
        self.assertLessEqual(result['coverage']['entries_enumerated'], 3)
        self.assertLessEqual(result['coverage']['entries_seen'], 3)
        self.assertLessEqual(result['coverage']['directories_seen'], 4)
        self.assertIn('total_entry_limit', {i['code'] for i in result['issues']})
        self.assertFalse(result['coverage']['complete'])

    def test_limits_validation(self):
        with self.assertRaises(ValueError):
            ScanLimits(max_files=0)
        with self.assertRaises(ValueError):
            ScanLimits(max_files=True)


if __name__ == '__main__':
    unittest.main()
