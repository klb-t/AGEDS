"""Real subprocess acceptance of the catalog CLI on the synthetic corpus."""
from __future__ import annotations

import csv
import json
import os
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from server.tests.corpus import catalog_files
from server.tests.test_catalog_scan import tree_state

REPO = Path(__file__).resolve().parents[2]


@pytest.fixture
def cli(tmp_path):
    corpus = catalog_files.build(tmp_path / 'corpus')
    data = tmp_path / 'data'
    env = {**os.environ, 'EW_DATA_DIR': str(data), 'EW_DB_PATH': str(data / 'evidence.db'), 'EW_STORE_DIR': str(data / 'store'),
           'EW_SCAN_ROOTS': ''}

    def run(*args, expect: int = 0) -> dict:
        process = subprocess.run([sys.executable, '-m', 'server.app.cli', *map(str, args)], cwd=REPO, env=env,
                                 capture_output=True, text=True, timeout=120)
        assert process.returncode == expect, (args, process.stdout, process.stderr)
        stream = process.stdout if expect == 0 else process.stderr
        return json.loads(stream.strip().splitlines()[-1])
    return corpus, tmp_path, run


def test_cli_end_to_end_without_touching_sources(cli):
    corpus, tmp, run = cli
    before = tree_state(tmp / 'corpus')
    a = run('catalog-collection-add', 'Phone A (synthetic)', '--device-or-account', 'phone-a', '--priority', '2')['id']
    b = run('catalog-collection-add', 'Drive B (synthetic)', '--parent', a)['id']
    updated = run('catalog-collection-update', a, '--notes', 'reviewed', '--clear', 'priority', '--actor', 'tester')
    assert updated['notes'] == 'reviewed' and updated['priority'] is None
    assert updated['history'][0]['details']['previous'] == {'notes': None, 'priority': 2}
    tree = run('catalog-collection-list')['collections']
    assert [(c['id'], c['tree_depth']) for c in tree] == [(a, 0), (b, 1)]

    scan_a = run('catalog-scan', corpus['roots']['phone_a'], '--collection', a, '--temp-dir', tmp / 'spool')
    assert scan_a['status'] == 'partial' and scan_a['source_bytes_written'] is False
    scan_b = run('catalog-scan', corpus['roots']['drive_b'], '--collection', b, '--temp-dir', tmp / 'spool',
                 '--max-archive-depth', '1')
    assert scan_b['coverage']['incomplete_reasons'] == {'container_not_expanded': 1, 'limit:max_archive_depth': 1}
    run('catalog-scan', corpus['roots']['drive_b'], '--collection', b, '--temp-dir', tmp / 'spool')
    assert len(run('catalog-scans', '--collection', b)['scans']) == 2

    coverage = run('catalog-coverage', '--collection', a, '--min-gap-days', '100')
    assert coverage['collections'][0]['totals']['gaps'] == 1
    duplicates = run('catalog-duplicates')['groups']
    assert duplicates[0]['entries'] == 3 and duplicates[0]['collections'] == 2
    group = run('catalog-duplicates', '--group-id', duplicates[0]['id'])
    assert {e['collection_id'] for e in group['catalog_entries'] if e['current_scan']} == {a, b}
    assert 'todo.txt' in {c['name'] for c in run('catalog-collisions')['collisions']}
    conflicts = run('catalog-entries', '--conflict', '--collection', a)['entries']
    assert {e['locator'] for e in conflicts} == set(corpus['type_conflicts'])

    out = tmp / 'exports'
    csv_result = run('catalog-export', out / 'catalog.csv', '--format', 'csv')
    with open(csv_result['output'], newline='', encoding='utf-8') as handle:
        rows = list(csv.DictReader(handle))
    assert len(rows) == csv_result['rows'] == csv_result['entries_selected']
    assert {r['locator'] for r in rows} >= {loc for _, loc, _ in corpus['shared_pdf']['locations']}
    xlsx = run('catalog-export', out / 'catalog.xlsx', '--format', 'xlsx', '--collection', a)
    workbook = load_workbook(xlsx['output'], read_only=True)
    assert sum(1 for _ in workbook['entries'].iter_rows()) == xlsx['rows'] + 1
    html = run('catalog-export', out / 'catalog.html', '--format', 'html')
    text = Path(html['output']).read_text(encoding='utf-8')
    assert '&lt;script&gt;alert(1).txt' in text and '<script' not in text
    again = run('catalog-export', out / 'catalog.csv', '--format', 'csv', expect=2)
    assert again['error'] == 'FileExistsError'
    inside = run('catalog-export', Path(corpus['roots']['phone_a']) / 'catalog.csv', '--format', 'csv', expect=2)
    assert 'outside' in inside['message']

    profile = run('catalog-describe', '--collection', a)
    assert profile['status'] == 'not_run' and profile['profile']['totals']['entries'] > 0
    hypothesis = run('catalog-describe', '--collection', a, '--provider', 'fake', '--output-dir', tmp / 'sidecars',
                     '--param', 'temperature=0')
    line = json.loads(Path(hypothesis['sidecar']).read_text().splitlines()[0])
    assert line['status'] == 'hypothesis' and line['format_compat'] == 'unknown' and line['params'] == {'temperature': '0'}
    assert run('catalog-describe', '--collection', a, '--provider', 'nonexistent', '--output-dir', tmp / 'sidecars',
               expect=2)['error'] == 'ValueError'
    decision = run('catalog-promote', line['annotation_id'], '--decision', 'reject', '--actor', 'reviewer', '--reason', 'too vague')
    assert decision['decision'] == 'rejected'
    assert run('catalog-mark', a, 'docs/report.pdf', '--priority', '1', '--note', 'key')['priority'] == 1
    deep = run('catalog-entries', '--q', 'deep.txt')['entries'][0]['id']
    ingested = run('catalog-ingest', deep, '--temp-dir', tmp / 'spool')
    assert ingested['ingested'] == 1
    assert tree_state(tmp / 'corpus') == before
    assert not list((tmp / 'spool').iterdir()), 'temporary copies are removed'


def test_cli_rejects_bad_arguments_without_creating_sources(cli):
    corpus, tmp, run = cli
    before = tree_state(tmp / 'corpus')
    process = subprocess.run([sys.executable, '-m', 'server.app.cli', 'catalog-scan', corpus['roots']['phone_a'],
                              '--collection', '1', '--max-files', '0'], cwd=REPO, capture_output=True, text=True,
                             env={**os.environ, 'EW_DATA_DIR': str(tmp / 'data')})
    assert process.returncode == 2 and 'positive integer' in process.stderr
    assert run('catalog-scan', corpus['roots']['phone_a'], '--collection', 99, expect=2)['message'] == 'collection does not exist'
    collection = run('catalog-collection-add', 'X')['id']
    assert run('catalog-scan', corpus['roots']['phone_a'], '--collection', collection, '--temp-dir',
               Path(corpus['roots']['phone_a']) / 'tmp', expect=2)['error'] == 'ValueError'
    assert run('catalog-collection-update', collection, '--parent', collection, expect=2)['error'] == 'ValueError'
    assert tree_state(tmp / 'corpus') == before
