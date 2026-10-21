"""Format writers: JSON, CSV and XLSX for tables, PDF from the HTML (adapter).

Exact values stay in ``report.json`` and the HTML. CSV text cells that a
spreadsheet would execute as a formula (leading ``= + - @`` TAB CR) get an
apostrophe prefix (guard ``apostrophe_prefix_v1``, recorded in the manifest).
XLSX string cells are forced to the string type so ``=...`` never becomes a
formula; characters XLSX cannot store are written as ``\\xNN`` text and counted.
XLSX needs openpyxl and PDF needs a Chromium executable (``EW_PDF_CHROMIUM``);
when missing the format is recorded as ``unavailable`` with the reason.
"""
from __future__ import annotations

import csv
import io
import json
import os
import re
import subprocess
import tempfile
from pathlib import Path

CSV_FORMULA_GUARD = 'apostrophe_prefix_v1'
_FORMULA_LEADS = ('=', '+', '-', '@', '\t', '\r')
PDF_ENV = 'EW_PDF_CHROMIUM'
PDF_TIMEOUT_SECONDS = 120
_SAFE_NAME = re.compile(r'[^A-Za-z0-9_.-]+')


def csv_guard(value):
    if isinstance(value, str) and value.startswith(_FORMULA_LEADS):
        return "'" + value
    return value


def _cell(value):
    if value is None:
        return ''
    if value is True or value is False:
        return 'true' if value else 'false'
    if isinstance(value, (int, float, str)):
        return value
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':'))


def table_blocks(document: dict):
    """``(section, block, pointer)`` for every table block, in document order."""
    for si, section in enumerate(document['sections']):
        for bi, block in enumerate(section['blocks']):
            if block['kind'] == 'table':
                yield section, block, f'/sections/{si}/blocks/{bi}'
    for bi, block in enumerate(document.get('method_view', [])):
        if block['kind'] == 'table':
            yield {'section': 'method'}, block, f'/method_view/{bi}'


def table_file_name(section: dict, block: dict, index: int) -> str:
    return f'tables/{index:03d}-{_SAFE_NAME.sub("_", section["section"] + "__" + block["id"])[:120]}.csv'


def write_csv_tables(document: dict, staging: Path) -> dict:
    files, guarded = [], 0
    (staging / 'tables').mkdir(mode=0o700)
    for index, (section, block, pointer) in enumerate(table_blocks(document)):
        name = table_file_name(section, block, index)
        buffer = io.StringIO(newline='')
        writer = csv.writer(buffer, lineterminator='\r\n')
        writer.writerow([column['key'] for column in block['columns']])
        for row in block['rows']:
            cells = []
            for column in block['columns']:
                value = _cell(row.get(column['key']))
                safe = csv_guard(value)
                guarded += safe is not value
                cells.append(safe)
            writer.writerow(cells)
        _write_new(staging / name, buffer.getvalue().encode('utf-8'))
        files.append({'path': name, 'pointer': pointer})
    return {'status': 'written', 'files': files, 'formula_guard': CSV_FORMULA_GUARD,
            'guarded_cells': guarded, 'encoding': 'utf-8', 'header': 'column keys'}


def xlsx_available() -> tuple[bool, str]:
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        return False, 'openpyxl is not installed (optional dependency, server/requirements-scanner.txt)'
    return True, 'openpyxl ' + __import__('openpyxl').__version__


def write_xlsx(document: dict, staging: Path) -> dict:
    available, detail = xlsx_available()
    if not available:
        return {'status': 'unavailable', 'reason': detail}
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    escaped = 0

    def put(sheet, row_index, column_index, value):
        nonlocal escaped
        value = _cell(value)
        if isinstance(value, str):
            clean, count = ILLEGAL_CHARACTERS_RE.subn(lambda match: '\\x%02x' % ord(match.group()), value)
            escaped += count
            cell = sheet.cell(row=row_index, column=column_index, value=clean)
            cell.data_type = 's'
        else:
            sheet.cell(row=row_index, column=column_index, value=value)

    book = Workbook()
    index_sheet = book.active
    index_sheet.title = 'index'
    for column, header in enumerate(('sheet', 'section', 'block', 'title', 'total_rows', 'truncated', 'pointer'), 1):
        put(index_sheet, 1, column, header)
    for index, (section, block, pointer) in enumerate(table_blocks(document)):
        name = f't{index:03d}'
        sheet = book.create_sheet(name)
        for column_index, column in enumerate(block['columns'], 1):
            put(sheet, 1, column_index, column['key'])
        for row_index, row in enumerate(block['rows'], 2):
            for column_index, column in enumerate(block['columns'], 1):
                put(sheet, row_index, column_index, row.get(column['key']))
        for column_index, value in enumerate((name, section['section'], block['id'], block.get('title'),
                                              block['total_rows'], block['truncated'], pointer), 1):
            put(index_sheet, index + 2, column_index, value)
    buffer = io.BytesIO()
    book.save(buffer)
    _write_new(staging / 'report.xlsx', buffer.getvalue())
    return {'status': 'written', 'files': [{'path': 'report.xlsx'}], 'library': detail,
            'string_cells': 'forced_string_type', 'illegal_characters_escaped': escaped}


def pdf_status() -> dict:
    executable = os.environ.get(PDF_ENV, '').strip()
    if not executable:
        return {'status': 'unavailable', 'reason': f'{PDF_ENV} is not set (path to a Chromium executable)'}
    if not (os.path.isfile(executable) and os.access(executable, os.X_OK)):
        return {'status': 'unavailable', 'reason': f'{PDF_ENV} does not name an executable file'}
    return {'status': 'available', 'executable': executable}


def write_pdf(staging: Path) -> dict:
    """Print ``report.html`` with headless Chromium; the HTML has no script/assets."""
    status = pdf_status()
    if status['status'] != 'available':
        return status
    executable = status['executable']
    try:
        version = subprocess.run([executable, '--version'], capture_output=True, text=True,
                                 timeout=30).stdout.strip() or 'unknown'
    except (OSError, subprocess.SubprocessError):
        version = 'unknown'
    target = staging / 'report.pdf'
    sandbox_disabled = hasattr(os, 'geteuid') and os.geteuid() == 0
    with tempfile.TemporaryDirectory(prefix='ageds-pdf-profile-') as profile:
        command = [executable, '--headless', '--disable-gpu', '--disable-dev-shm-usage', '--no-first-run',
                   '--no-default-browser-check', '--disable-extensions', '--disable-background-networking',
                   '--disable-sync', '--mute-audio', f'--user-data-dir={profile}', '--no-pdf-header-footer',
                   f'--print-to-pdf={target}', (staging / 'report.html').as_uri()]
        if sandbox_disabled:
            command.insert(1, '--no-sandbox')
        try:
            completed = subprocess.run(command, capture_output=True, timeout=PDF_TIMEOUT_SECONDS)
        except subprocess.TimeoutExpired:
            target.unlink(missing_ok=True)
            return {'status': 'failed', 'reason': f'chromium timed out after {PDF_TIMEOUT_SECONDS} s',
                    'chromium_version': version}
        except OSError as exc:
            return {'status': 'failed', 'reason': f'chromium could not start: {exc}', 'chromium_version': version}
    if completed.returncode != 0 or not target.is_file() or not target.read_bytes()[:5] == b'%PDF-':
        target.unlink(missing_ok=True)
        tail = completed.stderr.decode('utf-8', 'replace')[-500:]
        return {'status': 'failed', 'reason': f'chromium exit {completed.returncode}: {tail}', 'chromium_version': version}
    os.chmod(target, 0o600)
    return {'status': 'written', 'files': [{'path': 'report.pdf'}], 'adapter': 'chromium --headless --print-to-pdf',
            'chromium_version': version, 'sandbox_disabled_as_root': sandbox_disabled, 'source': 'report.html'}


def _write_new(path: Path, data: bytes) -> None:
    descriptor = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL | getattr(os, 'O_NOFOLLOW', 0), 0o600)
    with os.fdopen(descriptor, 'wb') as handle:
        handle.write(data)
        handle.flush()
        os.fsync(handle.fileno())


write_new = _write_new
